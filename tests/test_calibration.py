"""Calibration materials and checks (design 13.2, 13.4) on synthetic data: the 60-ticket set, the 30-ticket agreement subset, kappa, the blinding exercise and the recorded result."""
import json
import pathlib
import tempfile
import unittest

from planner_experiment import calibration as cal
from planner_experiment import review_layout
from planner_experiment import select_cases as sc
from tests.pe_world import KINDS, PARTITION_ARTIFACT, golden_raw
from planner_experiment import seed as seed_mod


def tickets(n=500):
    out = []
    for pool, wording in (("client_c", "problem reported on"), ("client_a", "fault noted against")):
        for i in range(n):
            out.append({"pool": pool, "tenant_id": pool, "ticket_id": f"INC{i:07d}", "text": f"{KINDS[i % len(KINDS)]} {wording} host{i}", "host": f"host{i}",
                        "opened_at": f"2026-{1 + i % 9:02d}-10 08:00:00"})
    return out


def goldens(n=18):
    return [{"ticket_id": f"golden-{i}", "text": f"calibration golden {i}", "host": f"gold{i}", "opened_at": "2026-05-01 10:00:00", "tenant_id": "poc_client"} for i in range(1, n + 1)]


class SetTests(unittest.TestCase):
    def setUp(self):
        self.dir = pathlib.Path(tempfile.mkdtemp())
        seed_mod.write_record(self.dir / "p.json", PARTITION_ARTIFACT, "calibration-selection")
        seed_mod.write_record(self.dir / "c.json", "d" * 40, "calibration-set")
        self.partition = seed_mod.load_record(self.dir / "p.json", "calibration-selection")["seed"]
        self.draw = seed_mod.load_record(self.dir / "c.json", "calibration-set")["seed"]

    def cases(self, **kw):
        return cal.draw_calibration(tickets(), goldens(), self.partition, self.draw, **kw)

    def test_sixty_cases_24_24_12_from_calibration_buckets_only(self):
        cases = self.cases()
        self.assertEqual(len(cases), 60)
        self.assertEqual({p: sum(1 for c in cases if c["pool"] == p) for p in ("client_c", "client_a", "golden")}, {"client_c": 24, "client_a": 24, "golden": 12})
        for c in cases:
            if c["source"] == "real":
                cid = sc.cluster_id({"pool": c["pool"], "tenant_id": c["tenant_id"], "text": c["ticket_text"], "host": c["host"]})
                self.assertIn(sc.bucket(self.partition, cid), sc.CALIBRATION_BUCKETS)

    def test_the_acceptance_buckets_are_never_touched(self):
        from planner_experiment import acceptance
        acc = acceptance.select_acceptance(tickets(), self.partition, seed_mod.derive_seed("c" * 40, "acceptance-selection"))["selected"]
        cases = self.cases()
        self.assertFalse({(c["pool"], c["ticket_id"]) for c in cases} & {(t["pool"], t["ticket_id"]) for t in acc})

    def test_development_tickets_and_their_goldens_are_excluded(self):
        first = self.cases()
        excluded = frozenset({(c["pool"], c["ticket_id"]) for c in first[:10]})
        clusters = frozenset(sc.cluster_id({**c, "text": c["ticket_text"]}) for c in first[:10] if c["pool"] != "golden")
        again = self.cases(exclude_tickets=excluded, exclude_clusters=clusters)
        self.assertFalse({(c["pool"], c["ticket_id"]) for c in again} & excluded)
        self.assertFalse({sc.cluster_id({**c, "text": c["ticket_text"]}) for c in again if c["pool"] != "golden"} & clusters)

    def test_it_is_deterministic_and_the_seed_matters(self):
        self.assertEqual([c["case_id"] for c in self.cases()], [c["case_id"] for c in self.cases()])
        other = cal.draw_calibration(tickets(), goldens(), self.partition, self.draw + 1)
        self.assertNotEqual([c["case_id"] for c in self.cases()], [c["case_id"] for c in other])

    def test_too_few_goldens_or_a_short_pool_is_an_error_not_a_smaller_set(self):
        with self.assertRaises(cal.CalibrationError):
            cal.draw_calibration(tickets(), goldens(8), self.partition, self.draw)
        with self.assertRaises(cal.CalibrationError):
            cal.draw_calibration(tickets(30), goldens(), self.partition, self.draw)

    def test_the_agreement_subset_and_the_blinding_sample_are_seeded_subsets(self):
        cases = self.cases()
        a = cal.assignments(cases, 7)
        self.assertEqual((len(a["all"]), len(a["agreement"]), len(a["blinding"])), (60, 30, 20))
        self.assertTrue(set(a["agreement"]) <= set(a["all"]) and set(a["blinding"]) <= set(a["all"]))
        self.assertEqual(a, cal.assignments(list(reversed(cases)), 7))
        self.assertNotEqual(a["agreement"], cal.assignments(cases, 8)["agreement"])


class KappaTests(unittest.TestCase):
    def test_known_values(self):
        self.assertEqual(cal.kappa(["a", "b"] * 5, ["a", "b"] * 5), 1.0)
        self.assertAlmostEqual(cal.kappa(["a"] * 5 + ["b"] * 5, ["a"] * 4 + ["b"] * 6), 0.8)
        self.assertAlmostEqual(cal.kappa(["a", "b", "a", "b"], ["b", "a", "b", "a"]), -1.0)
        self.assertEqual(cal.kappa(["a"] * 4, ["a"] * 4), 1.0)                      # no variation, full agreement
        with self.assertRaises(cal.CalibrationError):
            cal.kappa([], [])
        with self.assertRaises(cal.CalibrationError):
            cal.kappa(["a"], ["a", "b"])

    def labels(self, values):
        return {f"c{i}": {"determinable": v, "classification": "disk_full", "conflicting": "no"} for i, v in enumerate(values)}

    def test_the_gate_is_kappa_0_6_on_the_determinable_label(self):
        ids = [f"c{i}" for i in range(30)]
        base = ["determinable"] * 15 + ["not_determinable"] * 12 + ["indeterminate"] * 3
        good = list(base)
        good[0], good[15] = "not_determinable", "determinable"                      # two disagreements of thirty
        report = cal.agreement_report(self.labels(base), self.labels(good), ids)
        self.assertTrue(report["passed"])
        self.assertGreaterEqual(report["kappa_determinable"], 0.6)
        self.assertEqual(report["disagreements"], ["c0", "c15"])
        bad = ["not_determinable" if v == "determinable" else "determinable" for v in base]
        self.assertFalse(cal.agreement_report(self.labels(base), self.labels(bad), ids)["passed"])

    def test_incomplete_or_invalid_labels_are_refused(self):
        ids = [f"c{i}" for i in range(30)]
        full = self.labels(["determinable"] * 30)
        partial = {k: v for k, v in list(full.items())[:29]}
        with self.assertRaisesRegex(cal.CalibrationError, "not labelled"):
            cal.agreement_report(full, partial, ids)
        invalid = self.labels(["maybe"] + ["determinable"] * 29)
        with self.assertRaisesRegex(cal.CalibrationError, "invalid"):
            cal.agreement_report(full, invalid, ids)


def items(n=30):
    return [{"case_id": f"cal{i}", "evidence_packet": [f"Ticket {i}."], "baseline": {"causes": ["Disk quota too small", "Log growth"], "evidence": ["Quota exceeded on the volume"], "escalate": False},
             "candidate": {"causes": ["Log growth on the volume"], "evidence": ["Volume 98 percent full"], "escalate": i % 4 == 0}} for i in range(n)]


class BlindingAndResultTests(unittest.TestCase):
    def test_the_blinding_exercise_uses_the_blinded_layout_on_exactly_the_sample(self):
        sample = [f"cal{i}" for i in range(20)]
        packets, key = cal.blinding_materials(items(), sample, 3)
        self.assertEqual(len(packets), 20)
        self.assertTrue(review_layout.layout_is_identical(packets))
        with self.assertRaises(cal.CalibrationError):
            cal.blinding_materials(items(10), sample, 3)

    def test_the_result_passes_only_when_both_checks_pass_and_is_read_back_strictly(self):
        d = pathlib.Path(tempfile.mkdtemp())
        path = d / "r.json"
        good_a = {"passed": True, "kappa_determinable": 0.7, "n": 30}
        one = {"near_chance": True, "accuracy": 0.55, "n": 20}
        good_b = {"near_chance": True, "accuracy": 0.55, "reviewers": {"Vijay": one, "Sara": one}}
        self.assertFalse(cal.passed(path))                                           # absent
        cal.record_result(path, good_a, good_b, ["Vijay", "Sara"], "Ramanathan")
        self.assertTrue(cal.passed(path))
        cal.require_passed(path)
        bad_blinding = {"near_chance": False, "accuracy": 0.8, "reviewers": {"Vijay": {"near_chance": False, "accuracy": 0.8}, "Sara": one}}
        for a, b in (({"passed": False, "kappa_determinable": 0.4}, good_b), (good_a, bad_blinding)):
            cal.record_result(path, a, b, ["Vijay", "Sara"], "Ramanathan")
            self.assertFalse(cal.passed(path))
            with self.assertRaises(cal.CalibrationError):
                cal.require_passed(path)
        with self.assertRaises(cal.CalibrationError):
            cal.record_result(path, good_a, good_b, ["Vijay", "Sara"], " ")
        path.write_text(json.dumps({"passed": True, "agreement": {"passed": True, "kappa_determinable": 0.3}, "blinding": good_b, "calibration_tickets_only": True, "recorded_by": "x"}), encoding="utf-8")
        self.assertFalse(cal.passed(path))                                           # a hand-edited 'passed' with a failing kappa is not accepted

    def test_the_repository_has_no_passing_result_until_the_reviewers_have_done_the_checks(self):
        self.assertFalse(cal.passed())


class SheetTests(unittest.TestCase):
    def test_label_rows_hide_case_ids_and_carry_the_evidence(self):
        cases = [{"case_id": f"cal:client_c:INC{i}", "ticket_text": f"text {i}", "host": f"h{i}", "opened_at": "2026-01-01"} for i in range(3)]
        rows = cal.label_rows(cases, {"cal:client_c:INC0": ["related note"]}, [c["case_id"] for c in cases])
        self.assertEqual([r[0] for r in rows], ["C-001", "C-002", "C-003"])
        self.assertNotIn("INC", json.dumps(rows))
        self.assertIn("related note", rows[0][4])
        self.assertTrue(all(r[5] == "" for r in rows))                              # the label cells are empty

    def test_the_spreadsheet_round_trips_when_openpyxl_is_available(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl is not installed here")
        d = pathlib.Path(tempfile.mkdtemp())
        rows = cal.label_rows([{"case_id": "x1", "ticket_text": "t", "host": "h", "opened_at": "d"}], {}, ["x1"])
        path = d / "sheet.xlsx"
        cal.write_label_sheet(path, rows, "Vijay")
        wb = openpyxl.load_workbook(path)
        ws = wb["Labels"]
        ws["F2"], ws["G2"], ws["J2"] = "determinable", "disk_full", "no"
        wb.save(path)
        out = cal.read_label_sheet(path, {"C-001": "x1"})
        self.assertEqual(out["x1"]["determinable"], "determinable")
        self.assertEqual(out["x1"]["classification"], "disk_full")


class MaterialsTests(unittest.TestCase):
    def setUp(self):
        self.dir = pathlib.Path(tempfile.mkdtemp())
        seed_mod.write_record(self.dir / "p.json", PARTITION_ARTIFACT, "calibration-selection")
        part = seed_mod.load_record(self.dir / "p.json", "calibration-selection")["seed"]
        self.cases = cal.draw_calibration(tickets(), goldens(), part, 99)
        self.items = [{"case_id": c["case_id"], "evidence_packet": [c["ticket_text"]], "baseline": {"causes": ["Disk quota too small"], "evidence": ["Quota exceeded"], "escalate": False},
                       "candidate": {"causes": ["Log growth"], "evidence": ["Volume full"], "escalate": False}} for c in self.cases]
        self.out = self.dir / "materials"
        self.manifest = cal.build_label_materials(self.cases, self.out, ["Vijay", "Sara"], 5)

    def test_the_first_reviewer_gets_all_60_the_second_the_30_subset_and_no_system_output(self):
        self.assertEqual((self.manifest["files"]["Vijay"]["tickets"], self.manifest["files"]["Sara"]["tickets"]), (60, 30))
        for stem in ("vijay", "sara"):
            template = json.loads((self.out / f"labels_{stem}_template.json").read_text(encoding="utf-8"))
            text = json.dumps(template)
            self.assertNotIn("Log growth", text)                                                    # nothing from the candidate
            self.assertNotIn("Disk quota too small", text)
            self.assertNotIn("Earlier incident", text)                                              # no baseline-selected evidence either
            self.assertTrue(all(row[4] == "" for row in template["rows"]))
            self.assertTrue(all(row[5] == "" for row in template["rows"]))
        alias = json.loads((self.out / "alias_map_sara.json").read_text(encoding="utf-8"))
        self.assertEqual(set(alias.values()), set(self.manifest["plan"]["agreement"]))

    def test_the_blinding_packets_are_for_20_tickets_and_the_key_is_kept_apart(self):
        self.assertFalse((self.out / "blinding_packets.json").exists())                              # system output does not exist before labels
        ids = self.manifest["plan"]
        labels_a = {i: {"determinable": "determinable"} for i in ids["all"]}
        labels_b = {i: {"determinable": "determinable"} for i in ids["agreement"]}
        complete = cal.mark_labels_complete(self.manifest, labels_a, labels_b)
        cal.build_blinding_materials(self.items, self.out, complete, 5)
        packets = json.loads((self.out / "blinding_packets.json").read_text(encoding="utf-8"))["packets"]
        self.assertEqual(len(packets), 20)
        self.assertNotIn("baseline", json.dumps(packets))
        key = json.loads((self.out / "blinding_key.json").read_text(encoding="utf-8"))
        self.assertEqual(len(key), 20)

    def test_two_different_reviewers_are_required(self):
        with self.assertRaises(cal.CalibrationError):
            cal.build_label_materials(self.cases, self.dir / "x", ["Vijay", "vijay"], 5)

    def test_blinding_cannot_be_built_before_both_label_sheets_are_complete(self):
        with self.assertRaisesRegex(cal.CalibrationError, "after both"):
            cal.build_blinding_materials(self.items, self.out, self.manifest, 5)
        with self.assertRaisesRegex(cal.CalibrationError, "missing"):
            cal.mark_labels_complete(self.manifest, {}, {})

    def test_scoring_records_a_pass_only_for_good_agreement_and_chance_level_guessing(self):
        plan = self.manifest["plan"]
        labels_all = {i: {"determinable": "determinable" if n % 3 else "not_determinable", "classification": "x", "conflicting": "no"} for n, i in enumerate(plan["all"])}
        labels_b = {i: dict(labels_all[i]) for i in plan["agreement"]}
        complete = cal.mark_labels_complete(self.manifest, labels_all, labels_b)
        cal.build_blinding_materials(self.items, self.out, complete, 5)
        key = json.loads((self.out / "blinding_key.json").read_text(encoding="utf-8"))
        half = sorted(key)
        guesses = {alias: ("A" if n % 2 else "B") for n, alias in enumerate(half)}
        result_path = self.dir / "result.json"
        result = cal.score(labels_all, labels_b, plan, {"Vijay": guesses, "Sara": guesses}, key, ["Vijay", "Sara"], "Ramanathan", result_path)
        accuracy = result["blinding"]["accuracy"]
        self.assertEqual(result["passed"], accuracy <= 0.6)
        perfect = {alias: ("A" if k["A"] == "candidate" else "B") for alias, k in key.items()}
        failed = cal.score(labels_all, labels_b, plan, {"Vijay": perfect, "Sara": guesses}, key, ["Vijay", "Sara"], "Ramanathan", result_path)
        self.assertFalse(failed["passed"])
        self.assertFalse(cal.passed(result_path))
        with self.assertRaisesRegex(cal.CalibrationError, "both"):
            cal.score(labels_all, labels_b, plan, {"Vijay": guesses}, key, ["Vijay", "Sara"], "Ramanathan", result_path)


if __name__ == "__main__":
    unittest.main()
