"""
backtest_ingest.py — normalise client incident exports into a unified schema.

Reads all xlsx / xlsb files under data/rag/clients/<client_id>/ and outputs
data/backtest/incidents_normalised.json.

Run inside the container:
    python3 tools/backtest_ingest.py
"""
from __future__ import annotations
import json, os, re, pathlib, sys
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

ROOT = pathlib.Path(__file__).resolve().parent.parent
CLIENTS_DIR = ROOT / "data" / "rag" / "clients"
OUT_DIR = ROOT / "data" / "backtest"
OUT_FILE = OUT_DIR / "incidents_normalised.json"

# ---------------------------------------------------------------------------
# Column-name normalisation map  (alias → canonical)
# ---------------------------------------------------------------------------
COL_MAP = {
    # incident number
    "number": "incident_id",
    "number ": "incident_id",
    # priority
    "priority": "priority",
    # state / status
    "state": "status",
    "status": "status",
    # short description
    "short description": "short_description",
    "short_description": "short_description",
    # full description
    "description": "description",
    # assignment group
    "assignment group": "assignment_group",
    "assignment_group": "assignment_group",
    # category
    "category": "category",
    # subcategory
    "subcategory": "subcategory",
    # created
    "created": "created_at",
    "sys_created_on": "created_at",
    "opened_at": "created_at",
    # resolved / closed
    "resolved": "resolved_at",
    "closed_at": "resolved_at",
    "closed": "resolved_at",
    "resolved_at": "resolved_at",
    # resolution notes
    "resolution notes": "resolution_notes",
    "resolution_notes": "resolution_notes",
    "close_notes": "resolution_notes",
    # extras useful for backtesting
    "resolve time": "resolve_time_seconds",
    "reassignment count": "reassignment_count",
    "reassignment_count": "reassignment_count",
    "severity": "severity",
    "configuration item": "ci",
    "cmdb_ci": "ci",
    "configuration_item": "ci",
    "parent_incident": "parent_incident",
    "parent incident": "parent_incident",
    "major_incident_state": "major_incident",
    "is major incident?": "major_incident",
    "u_qs_major_incident": "major_incident",
    "towermap": "tower_map",
    "month": "month",
}


def _normalise_col(name: Any) -> str:
    return str(name or "").strip().lower().replace("\n", " ").replace("  ", " ")


def _safe_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _safe_date(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat()
    s = str(v).strip()
    # Excel serial date (float) — pyxlsb returns these
    if s.replace(".", "").isdigit():
        try:
            serial = float(s)
            # Excel epoch: 1899-12-30
            from datetime import timedelta
            base = datetime(1899, 12, 30, tzinfo=timezone.utc)
            return (base + timedelta(days=serial)).isoformat()
        except Exception:
            pass
    return s


def _priority_numeric(v: Any) -> Optional[int]:
    s = _safe_str(v)
    m = re.search(r"(\d)", s)
    return int(m.group(1)) if m else None


def _build_record(row_dict: Dict[str, Any], client_id: str, source_file: str) -> Optional[Dict[str, Any]]:
    inc_id = _safe_str(row_dict.get("incident_id", ""))
    if not inc_id or not inc_id.upper().startswith("INC"):
        return None

    short_desc = _safe_str(row_dict.get("short_description", ""))
    description = _safe_str(row_dict.get("description", ""))
    text = (short_desc + " " + description).strip()
    if not text:
        return None

    priority_raw = _safe_str(row_dict.get("priority", ""))
    priority_num = _priority_numeric(priority_raw)

    return {
        "incident_id": inc_id,
        "client_id": client_id,
        "source_file": source_file,
        "priority_raw": priority_raw,
        "priority": priority_num,
        "status": _safe_str(row_dict.get("status", "")),
        "short_description": short_desc,
        "description": description,
        "text": text,
        "assignment_group": _safe_str(row_dict.get("assignment_group", "")),
        "category": _safe_str(row_dict.get("category", "")),
        "subcategory": _safe_str(row_dict.get("subcategory", "")),
        "created_at": _safe_date(row_dict.get("created_at")),
        "resolved_at": _safe_date(row_dict.get("resolved_at")),
        "resolution_notes": _safe_str(row_dict.get("resolution_notes", "")),
        "reassignment_count": row_dict.get("reassignment_count"),
        "severity": _safe_str(row_dict.get("severity", "")),
        "ci": _safe_str(row_dict.get("ci", "")),
        "parent_incident": _safe_str(row_dict.get("parent_incident", "")),
        "major_incident": _safe_str(row_dict.get("major_incident", "")),
        "tower_map": _safe_str(row_dict.get("tower_map", "")),
    }


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

def _read_xlsx(path: pathlib.Path, client_id: str) -> List[Dict]:
    import openpyxl
    records = []
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        print(f"  SKIP {path.name}: {e}")
        return records

    for sh_name in wb.sheetnames:
        ws = wb[sh_name]
        headers: List[str] = []
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if not headers:
                raw_hdrs = [_normalise_col(c) for c in row]
                if not any(raw_hdrs):
                    continue
                # Map to canonical names; keep original for unmapped cols
                headers = [COL_MAP.get(h, h) for h in raw_hdrs]
                continue
            row_dict = dict(zip(headers, row))
            rec = _build_record(row_dict, client_id, f"{client_id}/{path.name}/{sh_name}")
            if rec:
                records.append(rec)
                row_count += 1
        if row_count:
            print(f"  {path.name} [{sh_name}]: {row_count} incidents")
    wb.close()
    return records


def _read_xlsb(path: pathlib.Path, client_id: str) -> List[Dict]:
    import pyxlsb
    records = []
    try:
        wb = pyxlsb.open_workbook(str(path))
    except Exception as e:
        print(f"  SKIP {path.name}: {e}")
        return records

    with wb:
        for sh_name in wb.sheets:
            with wb.get_sheet(sh_name) as ws:
                headers: List[str] = []
                row_count = 0
                for row in ws.rows():
                    vals = [c.v for c in row]
                    if not any(v for v in vals):
                        continue
                    if not headers:
                        raw_hdrs = [_normalise_col(c) for c in vals]
                        if not any(raw_hdrs):
                            continue
                        headers = [COL_MAP.get(h, h) for h in raw_hdrs]
                        continue
                    row_dict = dict(zip(headers, vals))
                    rec = _build_record(row_dict, client_id, f"{client_id}/{path.name}/{sh_name}")
                    if rec:
                        records.append(rec)
                        row_count += 1
                if row_count:
                    print(f"  {path.name} [{sh_name}]: {row_count} incidents")
    return records


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def ingest_all() -> List[Dict]:
    all_records: List[Dict] = []
    seen_ids: set = set()

    for client_dir in sorted(CLIENTS_DIR.iterdir()):
        if not client_dir.is_dir():
            continue
        client_id = client_dir.name
        print(f"\n[{client_id}]")
        before = len(all_records)

        for fpath in sorted(client_dir.iterdir()):
            if fpath.suffix.lower() == ".xlsx":
                recs = _read_xlsx(fpath, client_id)
            elif fpath.suffix.lower() == ".xlsb":
                recs = _read_xlsb(fpath, client_id)
            else:
                continue

            # Deduplicate within a client by incident_id
            for r in recs:
                key = f"{client_id}:{r['incident_id']}"
                if key not in seen_ids:
                    seen_ids.add(key)
                    all_records.append(r)

        added = len(all_records) - before
        print(f"  → {added} unique incidents from {client_id}")

    return all_records


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    print("=== Backtest Ingest ===")
    print(f"Source : {CLIENTS_DIR}")
    print(f"Output : {OUT_FILE}")

    records = ingest_all()

    # Summary by client
    from collections import Counter
    by_client = Counter(r["client_id"] for r in records)
    print(f"\n=== Summary ===")
    for cid, cnt in sorted(by_client.items()):
        print(f"  {cid}: {cnt:,} incidents")
    print(f"  TOTAL: {len(records):,} incidents")

    # Write output
    out = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total": len(records),
        "by_client": dict(by_client),
        "incidents": records,
    }
    OUT_FILE.write_text(json.dumps(out, indent=2, default=str), encoding="utf-8")
    print(f"\nWritten: {OUT_FILE}")


if __name__ == "__main__":
    main()
