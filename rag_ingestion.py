
import argparse, json, re, zipfile
from pathlib import Path
from typing import Any, List, Dict

try:
    import pandas as pd
except:
    pd = None
try:
    from docx import Document
except:
    Document = None
try:
    from pptx import Presentation
except:
    Presentation = None
try:
    from PIL import Image
except:
    Image = None
try:
    import pytesseract
except:
    pytesseract = None

SUPPORTED = {".csv",".xlsx",".xls",".json",".txt",".md",".docx",".pptx",".png",".jpg",".jpeg",".zip"}

def norm_text(v):
    return "" if v is None else str(v).strip()

def slug(v):
    v = norm_text(v).lower()
    return re.sub(r'[^a-z0-9]+','_',v).strip('_')

def write_json(path,payload):
    Path(path).write_text(json.dumps(payload,indent=2,ensure_ascii=False,default=str),encoding="utf-8")

def classify(name,text):
    t=(name+" "+text[:1000]).lower()
    if any(k in t for k in ["incident","ticket","sev","outage"]): return "incident"
    if any(k in t for k in ["vmware","rvtools","host","cluster","topology","network"]): return "topology"
    return "document"

def load_table(path):
    rows=[]
    if pd is None: raise RuntimeError("pandas not installed")
    if path.suffix.lower()==".csv":
        df=pd.read_csv(path)
        rows=df.to_dict("records")
    else:
        xls=pd.ExcelFile(path)
        for s in xls.sheet_names:
            df=pd.read_excel(path,sheet_name=s)
            for r in df.to_dict("records"):
                r["_sheet"]=s
                rows.append(r)
    return rows

def load_docx(path):
    if Document is None: raise RuntimeError("python-docx not installed")
    doc=Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def load_pptx(path):
    if Presentation is None: raise RuntimeError("python-pptx not installed")
    prs=Presentation(str(path))
    out=[]
    for i,slide in enumerate(prs.slides,1):
        texts=[]
        for shp in slide.shapes:
            if hasattr(shp,"text"):
                tx=shp.text.strip()
                if tx: texts.append(tx)
        if texts: out.append(f"Slide {i}: " + " | ".join(texts))
    return "\n".join(out)

def load_image(path):
    if Image is None: raise RuntimeError("Pillow not installed")
    txt=""
    if pytesseract:
        try:
            txt=pytesseract.image_to_string(Image.open(path))
        except:
            txt=""
    return f"Image file {path.name}\n{txt}"

def process_file(path):
    suf=path.suffix.lower()
    if path.name.startswith("~$"): raise RuntimeError("temp lock file skipped")
    if suf in [".csv",".xlsx",".xls"]:
        rows=load_table(path)
        out=[]
        for r in rows:
            txt=" | ".join(f"{k}:{v}" for k,v in r.items() if str(v)!="nan")
            out.append({"content":txt,"knowledge_type":classify(path.name,txt)})
        return out
    if suf in [".txt",".md",".json"]:
        txt=path.read_text(encoding="utf-8",errors="ignore")
        return [{"content":txt,"knowledge_type":classify(path.name,txt)}]
    if suf==".docx":
        txt=load_docx(path)
        return [{"content":txt,"knowledge_type":classify(path.name,txt)}]
    if suf==".pptx":
        txt=load_pptx(path)
        return [{"content":txt,"knowledge_type":classify(path.name,txt)}]
    if suf in [".png",".jpg",".jpeg"]:
        txt=load_image(path)
        return [{"content":txt,"knowledge_type":"topology"}]
    if suf==".zip":
        return [{"content":f"Archive file {path.name}","knowledge_type":"document"}]
    raise RuntimeError("unsupported")

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--tenant-id",default="client_a")
    ap.add_argument("--input-folder")
    ap.add_argument("input_file", nargs="?")
    args=ap.parse_args()

    files=[]
    if args.input_folder:
        for p in Path(args.input_folder).rglob("*"):
            if p.is_file() and p.suffix.lower() in SUPPORTED:
                files.append(p)
    elif args.input_file:
        files=[Path(args.input_file)]
    else:
        raise SystemExit("Provide file or --input-folder")

    incidents=[]; kb=[]; topo=[]; index=[]; manifest=[]; fails=[]
    collection=f"{slug(args.tenant_id)}_index"
    for f in files:
        try:
            recs=process_file(f)
            manifest.append({"file":str(f),"records":len(recs)})
            for rec in recs:
                item={"tenant_id":args.tenant_id,"collection_name":collection,
                      "source_file":str(f),"content":rec["content"][:8000],
                      "knowledge_type":rec["knowledge_type"]}
                index.append(item)
                if rec["knowledge_type"]=="incident": incidents.append(item)
                elif rec["knowledge_type"]=="topology": topo.append(item)
                else: kb.append(item)
        except Exception as e:
            fails.append({"file":str(f),"error":str(e)})

    out=Path.cwd()
    write_json(out/"tickets.json", incidents)
    write_json(out/"kb.json", kb)
    write_json(out/"topology.json", topo)
    write_json(out/"index_ready.json", index)
    write_json(out/"manifest.json", manifest)
    write_json(out/"failed_files.json", fails)
    print(json.dumps({
        "status":"ok","tenant_id":args.tenant_id,"collection_name":collection,
        "input_file_count":len(files),"ticket_count":len(incidents),
        "kb_count":len(kb),"topology_count":len(topo),
        "chunk_count":len(index),"failed_file_count":len(fails)
    },indent=2))
if __name__=="__main__":
    main()
