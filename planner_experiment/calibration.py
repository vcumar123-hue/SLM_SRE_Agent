"""
Calibration materials and checks for the labelling reviewers (design 13.2 and 13.4), done on **calibration tickets only**, before the acceptance freeze.

* a 60-ticket calibration set (24 PYPL + 24 PFG from the calibration buckets, 12 approved calibration goldens), clusters of the development subset excluded;
* an agreement exercise: one reviewer labels all 60, both reviewers label a seeded 30-ticket subset independently; Cohen's kappa on the *determinable* label must be at least 0.6 before labelling continues;
* a blinding exercise: on a seeded 20-ticket sample the reviewers see the baseline and candidate outputs in the blinded layout and guess which is the candidate; accuracy must be at most 60%;
* a recorded result (`docs/evaluation/calibration_result.json`) that the acceptance freeze requires to have passed.

The acceptance tickets are never involved here: the calibration buckets (0-39) and the acceptance buckets (40-79) of the same partition never share a cluster.
"""
from __future__ import annotations

import json
import random
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from planner_experiment import golden as golden_mod
from planner_experiment import review_layout
from planner_experiment import seed as seed_mod
from planner_experiment import select_cases as sc

CALIBRATION_QUOTAS = {"client_c": 24, "client_a": 24}
GOLDEN_COUNT = 12
SET_SIZE = 60
AGREEMENT_SUBSET = 30
BLINDING_SAMPLE = 20
MIN_KAPPA = 0.6
RESULT_PATH = Path(__file__).resolve().parent.parent / "docs" / "evaluation" / "calibration_result.json"
DETERMINABLE = ("determinable", "not_determinable", "indeterminate")


class CalibrationError(Exception):
    pass


def draw_calibration(tickets: Iterable[Dict[str, Any]], calibration_goldens: List[Dict[str, Any]], partition_seed: int, draw_seed: int,
                     exclude_tickets: frozenset = frozenset(), exclude_clusters: frozenset = frozenset()) -> List[Dict[str, Any]]:
    """60 calibration tickets as case dicts.

    Exact development tickets are excluded as defence in depth; their complete
    near-duplicate clusters are excluded through ``exclude_clusters``.
    """
    tickets = list(tickets)
    real = sc.select_calibration(tickets, partition_seed, draw_seed, quotas=dict(CALIBRATION_QUOTAS), exclude_tickets=exclude_tickets,
                                 exclude_clusters=exclude_clusters)
    if real["shortfall"]:
        raise CalibrationError(f"the calibration draw could not fill its quotas ({real['shortfall']})")
    rng = random.Random(draw_seed + 1)
    pool = [g for g in sorted(calibration_goldens, key=lambda g: str(g["ticket_id"])) if ("golden", str(g["ticket_id"])) not in exclude_tickets]
    if len(pool) < GOLDEN_COUNT:
        raise CalibrationError(f"only {len(pool)} calibration goldens are available, {GOLDEN_COUNT} are needed")
    goldens = rng.sample(pool, GOLDEN_COUNT)
    cases = [{"case_id": f"cal:{t['pool']}:{t['ticket_id']}", "pool": t["pool"], "tenant_id": t["tenant_id"], "ticket_id": t["ticket_id"], "ticket_text": t["text"],
              "host": t.get("host", ""), "opened_at": t.get("opened_at", ""), "source": "real"} for t in real["selected"]]
    cases += [{"case_id": f"cal:golden:{g['ticket_id']}", "pool": "golden", "tenant_id": g.get("tenant_id") or "poc_client", "ticket_id": g["ticket_id"], "ticket_text": g["text"],
               "host": g["host"], "opened_at": g["opened_at"], "source": "golden"} for g in goldens]
    cases.sort(key=lambda c: c["case_id"])
    if len({c["case_id"] for c in cases}) != SET_SIZE:
        raise CalibrationError("the calibration set must be 60 unique cases")
    return cases


def assignments(cases: Sequence[Dict[str, Any]], seed: int) -> Dict[str, List[str]]:
    """{'all': 60 ids for the primary reviewer, 'agreement': 30 ids both reviewers label, 'blinding': 20 ids for the blinding exercise}, all seeded."""
    ids = sorted(c["case_id"] for c in cases)
    rng = random.Random(seed)
    agreement = sorted(rng.sample(ids, AGREEMENT_SUBSET))
    blinding = sorted(rng.sample(ids, BLINDING_SAMPLE))
    return {"all": ids, "agreement": agreement, "blinding": blinding}


def kappa(a: Sequence[str], b: Sequence[str]) -> float:
    """Cohen's kappa for two raters' categorical labels over the same items."""
    if len(a) != len(b) or not a:
        raise CalibrationError("kappa needs two equally long, non-empty label lists")
    n = len(a)
    cats = sorted(set(a) | set(b))
    observed = sum(1 for x, y in zip(a, b) if x == y) / n
    ca, cb = Counter(a), Counter(b)
    expected = sum((ca[c] / n) * (cb[c] / n) for c in cats)
    if expected >= 1.0:
        return 1.0 if observed == 1.0 else 0.0
    return (observed - expected) / (1.0 - expected)


def agreement_report(labels_a: Dict[str, Dict[str, Any]], labels_b: Dict[str, Dict[str, Any]], subset_ids: Sequence[str]) -> Dict[str, Any]:
    """Compare two reviewers on the double-labelled subset. Gate: kappa on `determinable` >= 0.6. Also reported: agreement on conflicting and on the classification."""
    for name, labels in (("first", labels_a), ("second", labels_b)):
        missing = [i for i in subset_ids if i not in labels]
        if missing:
            raise CalibrationError(f"the {name} reviewer has not labelled {len(missing)} of the {len(subset_ids)} subset tickets")
        bad = [i for i in subset_ids if labels[i].get("determinable") not in DETERMINABLE]
        if bad:
            raise CalibrationError(f"the {name} reviewer's 'determinable' label is invalid on {bad[:3]}")
    da, db = [labels_a[i]["determinable"] for i in subset_ids], [labels_b[i]["determinable"] for i in subset_ids]
    k = kappa(da, db)
    pct = lambda field: round(100.0 * sum(1 for i in subset_ids if str(labels_a[i].get(field, "")).strip().lower() == str(labels_b[i].get(field, "")).strip().lower()) / len(subset_ids), 1)
    return {"n": len(subset_ids), "kappa_determinable": round(k, 3), "required": MIN_KAPPA, "passed": k >= MIN_KAPPA,
            "agreement_pct_classification": pct("classification"), "agreement_pct_conflicting": pct("conflicting"),
            "disagreements": [i for i, x, y in zip(subset_ids, da, db) if x != y]}


def blinding_materials(review_items: Sequence[Dict[str, Any]], blinding_ids: Sequence[str], seed: int):
    """Packets and key for the blinding exercise, from the executor's review items restricted to the sample."""
    chosen = [r for r in review_items if r["case_id"] in set(blinding_ids)]
    if len(chosen) != len(set(blinding_ids)):
        raise CalibrationError("review items are missing for part of the blinding sample")
    return review_layout.build_packets(chosen, seed)


def record_result(path: Path, agreement: Dict[str, Any], blinding: Dict[str, Any], reviewers: Sequence[str], recorded_by: str) -> Dict[str, Any]:
    if not str(recorded_by).strip():
        raise CalibrationError("a calibration result needs the name of who recorded it")
    result = {"recorded_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "recorded_by": recorded_by.strip(), "reviewers": list(reviewers), "agreement": agreement,
              "blinding": blinding, "passed": bool(agreement.get("passed") and blinding.get("near_chance")), "calibration_tickets_only": True}
    Path(path).write_text(json.dumps(result, indent=1) + "\n", encoding="utf-8")
    return result


def passed(path: Optional[Path] = None) -> bool:
    try:
        r = json.loads(Path(path or RESULT_PATH).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return False
    a, b = r.get("agreement") or {}, r.get("blinding") or {}
    reviewers = list(r.get("reviewers") or [])
    individual = b.get("reviewers") or {}
    both_blinded = len(reviewers) == 2 and set(individual) == set(reviewers) and all(
        v.get("near_chance") is True and float(v.get("accuracy", 1.0)) <= review_layout.BLINDING_MAX_ACCURACY for v in individual.values())
    return bool(r.get("passed") is True and a.get("passed") is True and float(a.get("kappa_determinable", -1)) >= MIN_KAPPA and b.get("near_chance") is True
                and both_blinded and r.get("calibration_tickets_only") is True and str(r.get("recorded_by") or "").strip())


def require_passed(path: Optional[Path] = None) -> None:
    if not passed(path):
        raise CalibrationError("the calibration checks (agreement kappa >= 0.6 and blinding accuracy <= 60%) have not been recorded as passed")


# -- sheets ---------------------------------------------------------------------------------------------------------------------------------

LABEL_COLUMNS = ["Ticket", "Ticket text", "Host", "Opened at", "Evidence the system could reach", "Determinable (determinable / not_determinable / indeterminate)",
                 "Correct classification", "Acceptable alternatives", "Likely root cause", "Conflicting evidence? (yes / no)", "Comment"]


def label_rows(cases: Sequence[Dict[str, Any]], evidence: Dict[str, List[str]], ids: Sequence[str]) -> List[List[Any]]:
    by_id = {c["case_id"]: c for c in cases}
    rows = []
    for alias, cid in enumerate(sorted(ids), 1):
        c = by_id[cid]
        rows.append([f"C-{alias:03d}", c["ticket_text"], c.get("host", ""), c.get("opened_at", ""), "\n".join(f"- {e}" for e in evidence.get(cid, [])), "", "", "", "", "", ""])
    return rows


def write_label_sheet(path: Path, rows: List[List[Any]], reviewer: str) -> None:
    try:
        import openpyxl
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise CalibrationError("openpyxl is needed to write the spreadsheet; use a Python that has it") from exc
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Labels"
    ws.append(LABEL_COLUMNS)
    for r in rows:
        ws.append(r)
    dv = DataValidation(type="list", formula1='"determinable,not_determinable,indeterminate"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(rows) + 1}")
    yn = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(yn)
    yn.add(f"J2:J{len(rows) + 1}")
    meta = wb.create_sheet("Reviewer")
    meta.append(["Reviewer name", reviewer])
    meta.append(["Date reviewed", ""])
    meta.append(["Instructions", "Label from the ticket and the evidence only. You have not seen any system output. Work alone; do not compare with the other reviewer."])
    wb.save(path)


def read_label_sheet(path: Path, alias_to_case: Dict[str, str]) -> Dict[str, Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise CalibrationError("openpyxl is needed to read the spreadsheet") from exc
    ws = openpyxl.load_workbook(path, data_only=True)["Labels"]
    out: Dict[str, Dict[str, Any]] = {}
    for row in list(ws.iter_rows(values_only=True))[1:]:
        alias = str(row[0] or "").strip()
        if alias not in alias_to_case:
            continue
        out[alias_to_case[alias]] = {"determinable": str(row[5] or "").strip().lower(), "classification": str(row[6] or "").strip(), "alternatives": str(row[7] or "").strip(),
                                     "root_cause": str(row[8] or "").strip(), "conflicting": str(row[9] or "").strip().lower(), "comment": str(row[10] or "").strip()}
    return out


# -- the package for the reviewers and the scoring of what they return ---------------------------------------------------------------------------

def build_label_materials(cases: Sequence[Dict[str, Any]], out_dir: Path, reviewers: Sequence[str], seed: int) -> Dict[str, Any]:
    """Write only the pre-output ground-truth labelling materials.

    These sheets deliberately contain the ticket and its own metadata only.  No
    baseline/candidate retrieval, evidence or output may be generated or shown
    until both completed label sheets have been returned.
    """
    if len(reviewers) != 2 or len({r.lower() for r in reviewers}) != 2:
        raise CalibrationError("two different reviewers are needed")
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    plan = assignments(cases, seed)
    manifest: Dict[str, Any] = {"reviewers": list(reviewers), "plan": plan, "files": {}}
    for reviewer, ids in ((reviewers[0], plan["all"]), (reviewers[1], plan["agreement"])):
        rows = label_rows(cases, {}, ids)
        alias_map = {row[0]: cid for row, cid in zip(rows, sorted(ids))}
        stem = reviewer.lower().replace(" ", "_")
        (out / f"labels_{stem}_template.json").write_text(json.dumps({"reviewer": reviewer, "columns": LABEL_COLUMNS, "rows": rows}, indent=1, ensure_ascii=False), encoding="utf-8")
        (out / f"alias_map_{stem}.json").write_text(json.dumps(alias_map, indent=1), encoding="utf-8")
        manifest["files"][reviewer] = {"template": f"labels_{stem}_template.json", "alias_map": f"alias_map_{stem}.json", "tickets": len(ids)}
        try:
            write_label_sheet(out / f"labels_{stem}.xlsx", rows, reviewer)
            manifest["files"][reviewer]["sheet"] = f"labels_{stem}.xlsx"
        except CalibrationError:
            manifest["files"][reviewer]["sheet"] = None                           # no openpyxl here: use the JSON template
    manifest["stage"] = "labels_only"
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    return manifest


def mark_labels_complete(manifest: Dict[str, Any], labels_a: Dict[str, Dict[str, Any]], labels_b: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """Validate that ground truth is complete before any system output exists."""
    plan = manifest.get("plan") or {}
    for name, labels, ids in (("first", labels_a, plan.get("all") or []), ("second", labels_b, plan.get("agreement") or [])):
        missing = [cid for cid in ids if cid not in labels]
        if missing:
            raise CalibrationError(f"the {name} reviewer's ground-truth sheet is missing {len(missing)} tickets")
        invalid = [cid for cid in ids if labels[cid].get("determinable") not in DETERMINABLE]
        if invalid:
            raise CalibrationError(f"the {name} reviewer has invalid determinability labels on {invalid[:3]}")
    updated = json.loads(json.dumps(manifest))
    updated["stage"] = "labels_complete"
    return updated


def build_blinding_materials(review_items: Sequence[Dict[str, Any]], out_dir: Path, manifest: Dict[str, Any], seed: int) -> Dict[str, Any]:
    """Add blinded system outputs only after ground-truth labels are complete.

    The caller is responsible for validating both returned label sheets before
    invoking this function.  The key remains custodian-only.
    """
    out = Path(out_dir)
    if manifest.get("stage") != "labels_complete":
        raise CalibrationError("blinding materials may be built only after both ground-truth label sheets are complete")
    packets, key = blinding_materials(review_items, manifest["plan"]["blinding"], seed + 2)
    (out / "blinding_packets.json").write_text(json.dumps({"instructions": "For each ticket, say which output (A or B) you believe came from the newer investigation system. Guess if unsure.", "packets": packets},
                                                       indent=1, ensure_ascii=False), encoding="utf-8")
    (out / "blinding_key.json").write_text(json.dumps(key, indent=1), encoding="utf-8")
    manifest["files"]["blinding"] = {"packets": "blinding_packets.json", "key": "blinding_key.json (not for the reviewers)", "tickets": len(packets)}
    manifest["stage"] = "blinding"
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    return manifest


def score(labels_a: Dict[str, Dict[str, Any]], labels_b: Dict[str, Dict[str, Any]], plan: Dict[str, List[str]], guesses_by_reviewer: Dict[str, Dict[str, str]], key: Dict[str, Dict[str, str]],
          reviewers: Sequence[str], recorded_by: str, result_path: Optional[Path] = None) -> Dict[str, Any]:
    """Agreement and two independent blinding checks; both reviewers must be at or below 60%."""
    agreement = agreement_report(labels_a, labels_b, plan["agreement"])
    if set(guesses_by_reviewer) != set(reviewers):
        raise CalibrationError("independent blinding guesses are required from both named reviewers")
    individual = {name: review_layout.blinding_accuracy(guesses_by_reviewer[name], key) for name in reviewers}
    blinding = {"reviewers": individual, "near_chance": all(v["near_chance"] for v in individual.values()),
                "accuracy": max(v["accuracy"] for v in individual.values())}
    return record_result(Path(result_path) if result_path else RESULT_PATH, agreement, blinding, reviewers, recorded_by)


def read_completed(path: Path, alias_map: Dict[str, str]) -> Dict[str, Dict[str, Any]]:
    """A reviewer's returned labels, from the .xlsx sheet or the completed .json template, keyed by case id."""
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        return read_label_sheet(p, alias_map)
    data = json.loads(p.read_text(encoding="utf-8"))
    out: Dict[str, Dict[str, Any]] = {}
    for row in data["rows"]:
        alias = str(row[0]).strip()
        if alias in alias_map:
            out[alias_map[alias]] = {"determinable": str(row[5] or "").strip().lower(), "classification": str(row[6] or "").strip(), "alternatives": str(row[7] or "").strip(),
                                     "root_cause": str(row[8] or "").strip(), "conflicting": str(row[9] or "").strip().lower(), "comment": str(row[10] or "").strip()}
    return out
